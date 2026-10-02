"""Parsers: um teste por formato presente na amostra (md, txt, css, svg, pdf, docx, pptx)."""
from __future__ import annotations

import docx
import pptx
import pytest

try:
    import pymupdf
except ImportError:  # pragma: no cover
    import fitz as pymupdf

from kentor_memoria.parsers import SUPPORTED_EXTENSIONS, parse_file


def test_formatos_suportados():
    assert {".md", ".txt", ".css", ".svg", ".pdf", ".docx", ".pptx"} <= SUPPORTED_EXTENSIONS


def test_markdown_frontmatter_headings_tabela(tmp_path):
    p = tmp_path / "doc.md"
    p.write_text(
        "---\ntitulo: Meu Doc\nversao: 2\n---\n# Título H1\n\nIntro.\n\n## Seção A\n\nTexto A com ![alt da figura](x.png).\n\n"
        "| col1 | col2 |\n|---|---|\n| a | b |\n\n### Sub A1\n\nTexto A1.\n\n```\n# não é heading\n```\n",
        encoding="utf-8",
    )
    d = parse_file(p)
    assert d.title == "Meu Doc" and d.metadata["frontmatter"]["versao"] == 2
    kinds = [(b.kind, b.heading_path) for b in d.blocks]
    assert ("table", ["Título H1", "Seção A"]) in kinds
    assert any(b.heading_path == ["Título H1", "Seção A", "Sub A1"] for b in d.blocks)
    assert any("alt da figura" in b.text for b in d.blocks)
    assert any(b.kind == "code" and "# não é heading" in b.text for b in d.blocks)


def test_txt_paragrafos(tmp_path):
    p = tmp_path / "nota.txt"
    p.write_text("Primeiro parágrafo\ncontinua.\n\nSegundo.\n", encoding="utf-8")
    d = parse_file(p)
    assert [b.text for b in d.blocks] == ["Primeiro parágrafo\ncontinua.", "Segundo."]


def test_css_secoes_por_comentario(tmp_path):
    p = tmp_path / "tokens.css"
    p.write_text(
        "/* ====\n   Tokens da Marca\n   ==== */\n:root {\n  /* === BRAND === */\n  --laranja: #FF8360;\n"
        "  /* === FONTS === */\n  --font: 'Sora';\n}\n",
        encoding="utf-8",
    )
    d = parse_file(p)
    assert d.title == "Tokens da Marca"
    brand = [b for b in d.blocks if b.heading_path == ["BRAND"]]
    assert brand and "#FF8360" in brand[0].text


def test_svg_textos_e_secoes(tmp_path):
    p = tmp_path / "org.svg"
    p.write_text(
        '<svg xmlns="http://www.w3.org/2000/svg"><title>Organograma</title>'
        "<!-- ===== Nível 1 ===== -->"
        '<text x="1" y="10">Vera</text><text x="50" y="10">Head de Vendas</text>'
        '<text x="1" y="30">Marta</text></svg>',
        encoding="utf-8",
    )
    d = parse_file(p)
    assert d.title == "Organograma"
    lvl = [b for b in d.blocks if b.heading_path == ["Nível 1"]][0]
    assert "Vera Head de Vendas" in lvl.text and "Marta" in lvl.text


def test_pdf_paginas_e_titulos(tmp_path):
    p = tmp_path / "deck.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Tabela de Preços", fontsize=24)
    page.insert_text((72, 110), "O pacote básico custa R$ 100 por mês.", fontsize=11)
    page = doc.new_page()
    page.insert_text((72, 72), "Garantia", fontsize=24)
    page.insert_text((72, 110), "Oferecemos 30 dias de garantia.", fontsize=11)
    page.insert_text((72, 130), "Sem letras miúdas.", fontsize=11)
    doc.set_metadata({"title": "Deck de Teste", "creationDate": "D:20260522000000"})
    doc.save(p)
    d = parse_file(p)
    assert d.title == "Deck de Teste" and d.metadata["page_count"] == 2 and d.metadata["created"] == "2026-05-22"
    garantia = [b for b in d.blocks if "30 dias" in b.text][0]
    assert garantia.page == 2 and garantia.heading_path == ["Garantia"]
    assert all("Tabela de Preços" != b.text for b in d.blocks)  # título não vira corpo


def test_docx_headings_paragrafos_tabelas_em_ordem(tmp_path):
    p = tmp_path / "t.docx"
    document = docx.Document()
    document.add_heading("Reunião", level=1)
    document.add_paragraph("Toda terça, 1 hora.")
    document.add_heading("Blocos", level=2)
    t = document.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Bloco", "Tempo"
    t.cell(1, 0).text, t.cell(1, 1).text = "Hooks", "20 min"
    document.save(p)
    d = parse_file(p)
    assert d.title == "Reunião"
    table = [b for b in d.blocks if b.kind == "table"][0]
    assert "Hooks | 20 min" in table.text and table.heading_path == ["Reunião", "Blocos"]
    assert d.blocks[0].text == "Toda terça, 1 hora." and d.blocks[0].heading_path == ["Reunião"]


def test_pptx_slides_titulo_tabela_notas(tmp_path):
    p = tmp_path / "s.pptx"
    prs = pptx.Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Canvas"
    s1.placeholders[1].text = "9 blocos de diagnóstico"
    s1.notes_slide.notes_text_frame.text = "Preencher numa manhã"
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Preços"
    rows = s2.shapes.add_table(2, 2, 0, 0, 3000000, 1000000).table
    rows.cell(0, 0).text, rows.cell(1, 0).text, rows.cell(1, 1).text = "Plano", "Básico", "R$ 10"
    prs.save(p)
    d = parse_file(p)
    assert [b.slide for b in d.blocks] == [1, 2]
    assert d.blocks[0].heading_path == ["Canvas"] and "Notas do apresentador: Preencher" in d.blocks[0].text
    assert "Básico | R$ 10" in d.blocks[1].text


def test_formato_nao_suportado(tmp_path):
    p = tmp_path / "x.xlsx"
    p.write_bytes(b"x")
    with pytest.raises(ValueError):
        parse_file(p)
