"""PPTX via python-pptx: um bloco por slide (texto, tabelas e notas)."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from ..models import Block, ParsedDocument
from ..utils import clean_text


def _shape_texts(shape) -> list[str]:
    out: list[str] = []
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        t = clean_text(shape.text_frame.text)
        if t:
            out.append(t)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [clean_text(c.text).replace("\n", " ") for c in row.cells]
            if any(cells):
                out.append(" | ".join(cells))
    # grupos de shapes
    if getattr(shape, "shapes", None) is not None and shape.shape_type == 6:  # GROUP
        for sub in shape.shapes:
            out.extend(_shape_texts(sub))
    return out


def parse_pptx(path: Path) -> ParsedDocument:
    prs = Presentation(str(path))
    blocks: list[Block] = []
    first_title: str | None = None
    for idx, slide in enumerate(prs.slides, start=1):
        title = None
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = clean_text(slide.shapes.title.text_frame.text) or None
        texts: list[str] = []
        for shape in slide.shapes:
            texts.extend(_shape_texts(shape))
        if not title:
            # sem placeholder de título: usa a primeira linha "curta" como título
            for t in texts:
                first = t.split("\n")[0].strip()
                if 3 <= len(first) <= 80:
                    title = first
                    break
        if slide.has_notes_slide:
            notes = clean_text(slide.notes_slide.notes_text_frame.text)
            if notes:
                texts.append(f"Notas do apresentador: {notes}")
        body = clean_text("\n".join(texts))
        if body:
            first_title = first_title or title
            blocks.append(Block(text=body, heading_path=[title] if title else [], slide=idx, kind="slide"))

    cp = prs.core_properties
    metadata = {
        "slide_count": len(prs.slides),
        "author": cp.author or None,
        "created": cp.created.date().isoformat() if cp.created else None,
        "modified": cp.modified.date().isoformat() if cp.modified else None,
    }
    return ParsedDocument(
        title=cp.title or first_title or path.stem,
        format="pptx",
        blocks=blocks,
        metadata={k: v for k, v in metadata.items() if v},
    )
